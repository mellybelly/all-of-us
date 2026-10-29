#!/usr/bin/env python3
"""Render the manuscript (outline.md) + references (references-seed.md) to a .docx
with headings, tables, bullets, inline bold/links/code, and clickable DOI hyperlinks."""
import re, sys
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.oxml.shared import OxmlElement, qn
from docx.oxml.ns import qn as nsqn

def add_hyperlink(p, url, text):
    part = p.part
    r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    h = OxmlElement('w:hyperlink'); h.set(qn('r:id'), r_id)
    r = OxmlElement('w:r'); rPr = OxmlElement('w:rPr')
    c = OxmlElement('w:color'); c.set(qn('w:val'), '1155CC'); rPr.append(c)
    u = OxmlElement('w:u'); u.set(qn('w:val'), 'single'); rPr.append(u)
    r.append(rPr)
    t = OxmlElement('w:t'); t.text = text; t.set(qn('xml:space'), 'preserve'); r.append(t)
    h.append(r); p._p.append(h)

TOKEN = re.compile(r'\*\*(.+?)\*\*|\[([^\]]+)\]\(([^)]+)\)|`([^`]+)`|\*(.+?)\*')
def render(p, text, size=11):
    pos = 0
    for m in TOKEN.finditer(text):
        if m.start() > pos:
            run = p.add_run(text[pos:m.start()]); run.font.size = Pt(size)
        if m.group(1) is not None:
            run = p.add_run(m.group(1)); run.bold = True; run.font.size = Pt(size)
        elif m.group(2) is not None:
            add_hyperlink(p, m.group(3), m.group(2))
        elif m.group(4) is not None:
            run = p.add_run(m.group(4)); run.font.name = 'Consolas'; run.font.size = Pt(size - 0.5)
        elif m.group(5) is not None:
            run = p.add_run(m.group(5)); run.italic = True; run.font.size = Pt(size)
        pos = m.end()
    if pos < len(text):
        run = p.add_run(text[pos:]); run.font.size = Pt(size)

doc = Document()
doc.styles['Normal'].font.name = 'Calibri'
doc.styles['Normal'].font.size = Pt(11)

def add_table(doc, rows):
    # rows: list of lists (header first), separator already dropped
    tbl = doc.add_table(rows=len(rows), cols=len(rows[0]))
    tbl.style = 'Light Grid Accent 1'
    for ri, row in enumerate(rows):
        for ci, cell in enumerate(row):
            para = tbl.cell(ri, ci).paragraphs[0]
            render(para, cell, size=9.5)
            if ri == 0:
                for run in para.runs:
                    run.bold = True

def split_row(line):
    cells = [c.strip() for c in line.strip().strip('|').split('|')]
    return cells

def emit(path):
    lines = [l.rstrip('\n') for l in open(path)]
    i = 0
    while i < len(lines):
        s = lines[i].strip()
        if not s or s == '---':
            i += 1; continue
        if s.startswith('|'):                      # table block
            block = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                block.append(lines[i]); i += 1
            rows = [split_row(b) for b in block]
            rows = [r for r in rows if not all(set(c) <= set('-: ') for c in r)]  # drop separator
            add_table(doc, rows); doc.add_paragraph()
            continue
        if s.startswith('> '):                     # blockquote callout
            p = doc.add_paragraph(); p.paragraph_format.left_indent = Inches(0.3)
            render(p, s[2:]);
            for r in p.runs: r.italic = True
            i += 1; continue
        if s.startswith('# '):
            doc.add_heading(re.sub(r'[*`]', '', s[2:]), 0)
        elif s.startswith('## '):
            doc.add_heading(re.sub(r'[*`]', '', s[3:]), 1)
        elif s.startswith('### '):
            doc.add_heading(re.sub(r'[*`]', '', s[4:]), 2)
        elif s.startswith('- '):
            render(doc.add_paragraph(style='List Bullet'), s[2:])
        elif re.match(r'^\d+\.\s', s):
            render(doc.add_paragraph(style='List Number'), re.sub(r'^\d+\.\s', '', s))
        else:
            render(doc.add_paragraph(), s)
        i += 1

emit('outline.md')
doc.add_page_break()
doc.add_heading('References (verified, with DOIs & links)', 0)
n = doc.add_paragraph()
rr = n.add_run('69 of 72 references carry a resolvable DOI (resolved via NCBI E-utilities); the rest link to their canonical source. Preprints are flagged to upgrade to journal versions before submission. Spot-check every DOI on import.')
rr.italic = True; rr.font.size = Pt(9.5)
emit('references-seed.md')

out = 'AoU_Roadmap_Refresh_Manuscript_Outline.docx'
doc.save(out)
print('wrote', out)
